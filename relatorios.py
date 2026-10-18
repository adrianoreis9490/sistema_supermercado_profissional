"""Exportação de relatórios de vendas em PDF e Excel."""
from pathlib import Path
from datetime import date
from banco import conectar
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import sys

BASE = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent
PASTA = BASE / "relatorios"
PASTA.mkdir(exist_ok=True)

def dados_mes(ano, mes):
    inicio = f"{ano:04d}-{mes:02d}-01"
    fim = f"{ano + (mes == 12):04d}-{1 if mes == 12 else mes + 1:02d}-01"
    with conectar() as conn:
        vendas = conn.execute("""
          SELECT v.id,v.data,v.total,v.pagamento,COALESCE(u.nome,'') AS funcionario
          FROM vendas v LEFT JOIN usuarios u ON u.id=v.funcionario_id
          WHERE v.data>=? AND v.data<? ORDER BY v.data
        """, (inicio, fim)).fetchall()
        resumo = conn.execute("""
          SELECT COUNT(*) AS quantidade, COALESCE(SUM(total),0) AS total
          FROM vendas WHERE data>=? AND data<?
        """, (inicio, fim)).fetchone()
        itens = conn.execute("""
          SELECT i.codigo_produto,i.nome_produto,SUM(i.quantidade) AS quantidade,
                 SUM(i.subtotal) AS valor
          FROM itens_venda i JOIN vendas v ON v.id=i.venda_id
          WHERE v.data>=? AND v.data<?
          GROUP BY i.codigo_produto,i.nome_produto ORDER BY quantidade DESC
        """, (inicio, fim)).fetchall()
    return vendas, resumo, itens

def exportar_excel(ano, mes):
    vendas, resumo, itens = dados_mes(ano, mes)
    caminho = PASTA / f"relatorio_{ano}_{mes:02d}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Vendas"
    ws.append(["ID", "Data", "Total (R$)", "Pagamento", "Funcionário"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="245A81")
    for v in vendas:
        ws.append([v["id"], v["data"], v["total"], v["pagamento"], v["funcionario"]])
    ws2 = wb.create_sheet("Produtos vendidos")
    ws2.append(["Código", "Produto", "Quantidade", "Valor (R$)"])
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="245A81")
    for i in itens:
        ws2.append([i["codigo_produto"], i["nome_produto"], i["quantidade"], i["valor"]])
    ws3 = wb.create_sheet("Resumo")
    ws3.append(["Mês", f"{mes:02d}/{ano}"])
    ws3.append(["Quantidade de vendas", resumo["quantidade"]])
    ws3.append(["Total vendido (R$)", resumo["total"]])
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        for col in sheet.columns:
            letter = col[0].column_letter
            sheet.column_dimensions[letter].width = min(max(max(len(str(c.value or "")) for c in col) + 2, 12), 36)
    wb.save(caminho)
    return caminho

def exportar_pdf(ano, mes):
    vendas, resumo, itens = dados_mes(ano, mes)
    caminho = PASTA / f"relatorio_{ano}_{mes:02d}.pdf"
    pdf = canvas.Canvas(str(caminho), pagesize=A4)
    _, altura = A4
    y = altura - 50
    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawString(45, y, f"Relatório mensal - {mes:02d}/{ano}")
    y -= 30
    pdf.setFont("Helvetica", 11)
    pdf.drawString(45, y, f"Quantidade de vendas: {resumo['quantidade']}")
    y -= 18
    pdf.drawString(45, y, f"Total vendido: R$ {resumo['total']:.2f}")
    y -= 30
    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(45, y, "Produtos mais vendidos")
    y -= 20
    pdf.setFont("Helvetica", 9)
    for item in itens:
        if y < 55:
            pdf.showPage()
            y = altura - 50
        pdf.drawString(45, y, f"{item['codigo_produto']} - {item['nome_produto'][:35]}")
        pdf.drawRightString(420, y, f"Qtd.: {item['quantidade']}")
        pdf.drawRightString(550, y, f"R$ {item['valor']:.2f}")
        y -= 15
    pdf.save()
    return caminho
